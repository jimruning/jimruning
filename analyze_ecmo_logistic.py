# -*- coding: utf-8 -*-

from __future__ import annotations

import math
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
import statsmodels.formula.api as smf
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from scipy.stats import chi2, chi2_contingency, fisher_exact, mannwhitneyu, shapiro, ttest_ind


BASE_DIR = Path(__file__).resolve().parent
DATA_FILE = BASE_DIR / "先心ECMO临床资料.xlsx"
RESULT_DIR = BASE_DIR / "results"
RESULT_DIR.mkdir(exist_ok=True)

PROCESSED_FILE = RESULT_DIR / "analysis_dataset_sheet2.csv"
EXCEL_FILE = RESULT_DIR / "先心ECMO_Logistic回归分析.xlsx"
SUMMARY_FILE = RESULT_DIR / "分析说明.md"
RESULT_NARRATIVE_FILE = RESULT_DIR / "回归分析结果描述.md"
PUBLISH_EXCEL_FILE = RESULT_DIR / "先心ECMO_Logistic回归分析_发表格式.xlsx"
FOREST_DIR = RESULT_DIR / "forest_plots"
FOREST_DIR.mkdir(exist_ok=True)

OUTCOMES = {
    "撤机成功": "撤机",
    "出院成功": "出院",
}

MISSING_TOKENS = {"", "nan", "none", "null", "缺", "缺失", "?", "？", "不详", "未记录"}
QUANTITATIVE_SOURCE_COLUMNS = {
    "年龄（m）",
    "VIS",
    "最高总胆红素",
    "最高lac",
    "12小时lac清除率",
    "RBC量u",
    "血小板u",
    "血浆量",
    "Fib量",
    "时长h",
    "体重kg",
    "亚里士多德评分",
}
CATEGORICAL_SOURCE_COLUMNS = {"性别", "出血", "血栓"}
LABEL_MAP = {
    "年龄（m）": "年龄（月）",
    "VIS": "VIS评分",
    "最高总胆红素": "最高总胆红素",
    "最高lac": "最高乳酸",
    "12小时lac清除率": "12小时乳酸清除率",
    "RBC量u": "红细胞输注量（U）",
    "血小板u": "血小板输注量（U）",
    "血浆量": "血浆输注量（mL）",
    "Fib量": "纤维蛋白原用量",
    "时长h": "ECMO时长（h）",
    "体重kg": "体重（kg）",
    "性别": "性别",
    "出血": "出血",
    "血栓": "血栓",
    "撤机成功": "撤机成功",
    "出院成功": "出院成功",
}

plt.rcParams["font.sans-serif"] = ["PingFang SC", "Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False


def normalize_text(value) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip()
    if not text:
        return None
    return text


def parse_binary_outcome(series: pd.Series) -> pd.Series:
    def convert(value):
        text = normalize_text(value)
        if text is None:
            return np.nan
        compact = re.sub(r"\s+", "", text)
        if "是" in compact:
            return 1
        if "否" in compact:
            return 0
        return np.nan

    return series.map(convert).astype("float")


def parse_numeric(series: pd.Series) -> pd.Series:
    cleaned = (
        series.astype(str)
        .str.strip()
        .replace(
            {
                "nan": np.nan,
                "NaN": np.nan,
                "None": np.nan,
                "不计": np.nan,
                "不详": np.nan,
                "？": np.nan,
                "?": np.nan,
                "缺": np.nan,
                "缺失": np.nan,
            }
        )
    )
    return pd.to_numeric(cleaned, errors="coerce")


def clean_categorical(series: pd.Series, source_name: str) -> pd.Series:
    def convert(value):
        text = normalize_text(value)
        if text is None:
            return np.nan
        compact = re.sub(r"\s+", "", text)
        if compact.lower() in MISSING_TOKENS:
            return np.nan
        if source_name == "性别":
            if compact in {"男", "male", "Male", "M"}:
                return "男"
            if compact in {"女", "female", "Female", "F"}:
                return "女"
            return np.nan
        if source_name in {"出血", "血栓"}:
            if compact in {"有", "是", "阳性"}:
                return "有"
            if compact in {"无", "否", "阴性"}:
                return "无"
            return np.nan
        return compact

    return series.map(convert).astype("object")


def choose_reference(series: pd.Series) -> str:
    values = set(series.dropna().astype(str).tolist())
    if "无" in values:
        return "无"
    counts = series.dropna().value_counts()
    if not counts.empty:
        return str(counts.index[0])
    return str(series.value_counts().index[0])


def make_tertiles(series: pd.Series) -> tuple[pd.Series, str]:
    non_missing = series.dropna()
    ranked = non_missing.rank(method="first")
    tertiles = pd.qcut(ranked, 3, labels=["Q1", "Q2", "Q3"])
    out = pd.Series(index=series.index, dtype="object")
    out.loc[non_missing.index] = tertiles.astype(str)
    return out, "Q1"


def is_normal_distribution(series: pd.Series) -> bool:
    non_missing = series.dropna()
    if len(non_missing) < 3:
        return True
    sample = non_missing
    if len(sample) > 5000:
        sample = sample.sample(5000, random_state=42)
    try:
        p_value = shapiro(sample)[1]
    except Exception:
        return False
    return bool(p_value >= 0.05)


def sanitize_name(name: str) -> str:
    ascii_name = re.sub(r"\W+", "_", name.strip())
    ascii_name = re.sub(r"_+", "_", ascii_name).strip("_")
    if not ascii_name:
        ascii_name = "var"
    if ascii_name[0].isdigit():
        ascii_name = f"v_{ascii_name}"
    return ascii_name


def safe_exp(value: float) -> float:
    if pd.isna(value):
        return np.nan
    if value > 700:
        return np.inf
    if value < -700:
        return 0.0
    return float(math.exp(value))


def fmt_num(value: float, digits: int = 3) -> str:
    if pd.isna(value):
        return ""
    if np.isinf(value):
        return "Inf"
    return f"{value:.{digits}f}"


def label_name(name: str) -> str:
    return LABEL_MAP.get(name, name)


def fmt_pct(n: int, total: int) -> str:
    if total <= 0:
        return "0 (0.0%)"
    return f"{n} ({n / total * 100:.1f}%)"


def fmt_continuous(series: pd.Series, normal: bool) -> str:
    clean = series.dropna()
    if clean.empty:
        return ""
    if normal:
        return f"{clean.mean():.2f} ({clean.std(ddof=1):.2f})"
    q1 = clean.quantile(0.25)
    q3 = clean.quantile(0.75)
    return f"{clean.median():.2f} ({q1:.2f}, {q3:.2f})"


def build_analysis_dataset() -> tuple[pd.DataFrame, dict[str, dict]]:
    df = pd.read_excel(DATA_FILE, sheet_name="Sheet2").copy()
    df.columns = [str(c).strip() for c in df.columns]

    for label, source_col in OUTCOMES.items():
        df[label] = parse_binary_outcome(df[source_col])

    predictor_columns = [
        col
        for col in df.columns
        if col not in {"姓名", "撤机", "出院", "撤机成功", "出院成功"}
        and not col.startswith("Unnamed")
    ]

    analysis_df = pd.DataFrame(index=df.index)
    metadata: dict[str, dict] = {}

    for col in predictor_columns:
        raw = df[col]
        safe_name = sanitize_name(col)
        while safe_name in analysis_df.columns:
            safe_name = f"{safe_name}_x"

        numeric = parse_numeric(raw)
        unique_numeric = numeric.dropna().nunique()

        if col in QUANTITATIVE_SOURCE_COLUMNS:
            if numeric.notna().sum() == 0 or unique_numeric < 3:
                continue
            if is_normal_distribution(numeric):
                analysis_df[safe_name] = numeric
                metadata[safe_name] = {
                    "source": col,
                    "type": "continuous",
                    "reference": None,
                    "transform": "原始连续变量",
                }
            else:
                tertile_series, reference = make_tertiles(numeric)
                analysis_df[safe_name] = tertile_series
                metadata[safe_name] = {
                    "source": col,
                    "type": "categorical",
                    "reference": reference,
                    "transform": "非正态，按三分位数组，缺失剔除",
                }
        elif col in CATEGORICAL_SOURCE_COLUMNS:
            cleaned = clean_categorical(raw, col)
            analysis_df[safe_name] = cleaned
            metadata[safe_name] = {
                "source": col,
                "type": "categorical",
                "reference": choose_reference(cleaned),
                "transform": "分类变量，缺失剔除",
            }
        else:
            cleaned = clean_categorical(raw, col)
            if cleaned.nunique(dropna=False) <= 1:
                continue
            analysis_df[safe_name] = cleaned
            metadata[safe_name] = {
                "source": col,
                "type": "categorical",
                "reference": choose_reference(cleaned),
                "transform": "分类变量，缺失剔除",
            }

    for outcome in OUTCOMES:
        analysis_df[outcome] = df[outcome]

    return analysis_df, metadata


def build_descriptive_dataset() -> tuple[pd.DataFrame, dict[str, dict]]:
    df = pd.read_excel(DATA_FILE, sheet_name="Sheet2").copy()
    df.columns = [str(c).strip() for c in df.columns]

    desc_df = pd.DataFrame(index=df.index)
    desc_meta: dict[str, dict] = {}

    for label, source_col in OUTCOMES.items():
        desc_df[label] = parse_binary_outcome(df[source_col])

    predictor_columns = [
        col
        for col in df.columns
        if col not in {"姓名", "撤机", "出院", "撤机成功", "出院成功"}
        and not col.startswith("Unnamed")
    ]

    for col in predictor_columns:
        if col in QUANTITATIVE_SOURCE_COLUMNS:
            numeric = parse_numeric(df[col])
            if numeric.notna().sum() == 0 or numeric.dropna().nunique() < 2:
                continue
            desc_df[col] = numeric
            desc_meta[col] = {
                "source": col,
                "type": "continuous",
                "normal": is_normal_distribution(numeric),
            }
        else:
            cleaned = clean_categorical(df[col], col)
            if cleaned.dropna().nunique() < 2:
                continue
            desc_df[col] = cleaned
            desc_meta[col] = {
                "source": col,
                "type": "categorical",
            }

    return desc_df, desc_meta


def build_group_comparison_table(data: pd.DataFrame, outcome: str, metadata: dict[str, dict]) -> pd.DataFrame:
    rows = []
    outcome_df = data[[outcome]].copy()
    outcome_df = outcome_df[outcome_df[outcome].notna()].copy()
    group0_label = "否"
    group1_label = "是"

    for predictor, meta in metadata.items():
        sub = data[[outcome, predictor]].copy()
        sub = sub[sub[outcome].notna() & sub[predictor].notna()].copy()
        if sub.empty or sub[outcome].nunique() < 2:
            continue

        g0 = sub.loc[sub[outcome] == 0, predictor]
        g1 = sub.loc[sub[outcome] == 1, predictor]
        total = sub[predictor]

        if meta["type"] == "continuous":
            normal = bool(meta.get("normal", True))
            stat_name = "t值" if normal else "Z值"
            try:
                if normal:
                    stat, p_value = ttest_ind(g1, g0, equal_var=False, nan_policy="omit")
                else:
                    stat, p_value = mannwhitneyu(g1, g0, alternative="two-sided")
            except Exception:
                stat, p_value = np.nan, np.nan

            rows.append(
                {
                    "变量": label_name(meta["source"]),
                    "水平/描述": "均数(标准差)" if normal else "中位数(P25, P75)",
                    f"{outcome}={group0_label}": fmt_continuous(g0, normal),
                    f"{outcome}={group1_label}": fmt_continuous(g1, normal),
                    "总计": fmt_continuous(total, normal),
                    "检验方法": "Welch t检验" if normal else "Mann-Whitney U检验",
                    "统计量": fmt_num(stat),
                    "P值": "<0.001" if pd.notna(p_value) and p_value < 0.001 else ("" if pd.isna(p_value) else f"{p_value:.3f}"),
                    "纳入样本量": len(sub),
                }
            )
            continue

        contingency = pd.crosstab(sub[predictor], sub[outcome])
        test_name = "卡方检验"
        stat_value = np.nan
        p_value = np.nan
        try:
            if contingency.shape == (2, 2):
                chi2_stat, chi2_p, _, expected = chi2_contingency(contingency)
                if (expected < 5).any():
                    stat_value, p_value = fisher_exact(contingency.to_numpy())
                    test_name = "Fisher确切检验"
                else:
                    stat_value, p_value = chi2_stat, chi2_p
            else:
                stat_value, p_value, _, _ = chi2_contingency(contingency)
        except Exception:
            stat_value, p_value = np.nan, np.nan

        levels = list(pd.Index(total.dropna().unique()))
        if meta.get("reference") in levels:
            levels = [meta["reference"]] + [lvl for lvl in levels if lvl != meta["reference"]]
        else:
            levels = sorted(levels, key=str)

        total0 = int((sub[outcome] == 0).sum())
        total1 = int((sub[outcome] == 1).sum())
        total_all = len(sub)

        for idx, level in enumerate(levels):
            n0 = int(((sub[outcome] == 0) & (sub[predictor] == level)).sum())
            n1 = int(((sub[outcome] == 1) & (sub[predictor] == level)).sum())
            n_all = int((sub[predictor] == level).sum())
            rows.append(
                {
                    "变量": label_name(meta["source"]),
                    "水平/描述": str(level),
                    f"{outcome}={group0_label}": fmt_pct(n0, total0),
                    f"{outcome}={group1_label}": fmt_pct(n1, total1),
                    "总计": fmt_pct(n_all, total_all),
                    "检验方法": test_name if idx == 0 else "",
                    "统计量": fmt_num(stat_value) if idx == 0 else "",
                    "P值": ("<0.001" if pd.notna(p_value) and p_value < 0.001 else ("" if pd.isna(p_value) else f"{p_value:.3f}")) if idx == 0 else "",
                    "纳入样本量": total_all if idx == 0 else "",
                }
            )

    return pd.DataFrame(rows)


def build_formula(outcome: str, predictor: str, meta: dict) -> str:
    if meta["type"] == "continuous":
        return f'Q("{outcome}") ~ Q("{predictor}")'
    reference = meta["reference"].replace('"', '\\"')
    return f'Q("{outcome}") ~ C(Q("{predictor}"), Treatment(reference="{reference}"))'


def fit_model(formula: str, data: pd.DataFrame):
    return smf.glm(formula=formula, data=data, family=sm.families.Binomial()).fit()


def likelihood_ratio_p(full_result, reduced_result) -> float:
    lr_stat = 2 * (full_result.llf - reduced_result.llf)
    df_diff = int(round(full_result.df_model - reduced_result.df_model))
    if df_diff <= 0:
        return np.nan
    return float(chi2.sf(lr_stat, df_diff))


def prepare_model_data(data: pd.DataFrame, outcome: str, predictor: str, meta: dict) -> pd.DataFrame:
    cols = [outcome, predictor]
    model_df = data[cols].copy()
    model_df = model_df[model_df[outcome].notna()].copy()
    model_df = model_df[model_df[predictor].notna()].copy()
    return model_df


def parse_term(term: str, predictor: str, meta: dict) -> tuple[str, str]:
    source = label_name(meta["source"])
    if meta["type"] == "continuous":
        return source, "每增加1单位"
    match = re.search(r'\[T\.(.*)\]$', term)
    level = match.group(1) if match else term
    level = {"Q1": "Q1（低三分位组）", "Q2": "Q2（中三分位组）", "Q3": "Q3（高三分位组）", "缺失": "缺失组"}.get(level, level)
    reference = {"Q1": "Q1（低三分位组）", "Q2": "Q2（中三分位组）", "Q3": "Q3（高三分位组）", "缺失": "缺失组"}.get(meta["reference"], meta["reference"])
    return source, f"{level} vs {reference}"


def run_univariate(data: pd.DataFrame, outcome: str, metadata: dict[str, dict]) -> tuple[pd.DataFrame, pd.DataFrame]:
    coef_rows = []
    variable_rows = []

    for predictor, meta in metadata.items():
        model_df = prepare_model_data(data, outcome, predictor, meta)
        if model_df.empty or model_df[outcome].nunique() < 2:
            continue
        if meta["type"] == "categorical" and model_df[predictor].nunique() < 2:
            continue
        if meta["type"] == "continuous" and model_df[predictor].nunique() < 2:
            continue

        formula = build_formula(outcome, predictor, meta)
        try:
            result = fit_model(formula, model_df)
            null_result = fit_model(f'Q("{outcome}") ~ 1', model_df)
        except Exception:
            continue

        overall_p = likelihood_ratio_p(result, null_result)
        variable_rows.append(
            {
                "predictor": predictor,
                "变量": label_name(meta["source"]),
                "变量类型": meta["transform"],
                "整体P值": overall_p,
                "样本量": len(model_df),
            }
        )

        conf_int = result.conf_int()
        for term in result.params.index:
            if term == "Intercept":
                continue
            coef = float(result.params[term])
            se = float(result.bse[term])
            p_value = float(result.pvalues[term])
            lower = float(conf_int.loc[term, 0])
            upper = float(conf_int.loc[term, 1])
            variable_name, comparison = parse_term(term, predictor, meta)
            coef_rows.append(
                {
                    "变量": variable_name,
                    "比较": comparison,
                    "变量类型": meta["transform"],
                    "回归系数β": coef,
                    "标准误SE": se,
                    "OR": safe_exp(coef),
                    "OR 95%CI下限": safe_exp(lower),
                    "OR 95%CI上限": safe_exp(upper),
                    "P值": p_value,
                    "整体P值": overall_p,
                    "样本量": len(model_df),
                }
            )

    coef_df = pd.DataFrame(coef_rows).sort_values(["整体P值", "P值", "变量"], na_position="last")
    variable_df = pd.DataFrame(variable_rows).sort_values(["整体P值", "变量"], na_position="last")
    return coef_df.reset_index(drop=True), variable_df.reset_index(drop=True)


def fit_multivariable_formula(data: pd.DataFrame, outcome: str, predictors: list[str], metadata: dict[str, dict]):
    terms = []
    for predictor in predictors:
        meta = metadata[predictor]
        if meta["type"] == "continuous":
            terms.append(f'Q("{predictor}")')
        else:
            ref = meta["reference"].replace('"', '\\"')
            terms.append(f'C(Q("{predictor}"), Treatment(reference="{ref}"))')
    formula = f'Q("{outcome}") ~ ' + " + ".join(terms)
    model_df = data[[outcome] + predictors].copy()
    model_df = model_df[model_df[outcome].notna()].copy()
    for predictor in predictors:
        model_df = model_df[model_df[predictor].notna()].copy()
    return fit_model(formula, model_df), model_df


def run_multivariate(
    data: pd.DataFrame,
    outcome: str,
    metadata: dict[str, dict],
    univariate_variable_df: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, list[str]]:
    selected = univariate_variable_df.loc[univariate_variable_df["整体P值"] < 0.05, "predictor"].tolist()
    if not selected:
        return pd.DataFrame(), pd.DataFrame(), []

    while len(selected) > 0:
        try:
            full_result, model_df = fit_multivariable_formula(data, outcome, selected, metadata)
        except Exception:
            if len(selected) == 1:
                return pd.DataFrame(), pd.DataFrame(), []
            selected = selected[:-1]
            continue

        worst_predictor = None
        worst_p = -1.0
        for predictor in list(selected):
            reduced_predictors = [p for p in selected if p != predictor]
            if not reduced_predictors:
                continue
            try:
                reduced_result, _ = fit_multivariable_formula(model_df, outcome, reduced_predictors, metadata)
                p_value = likelihood_ratio_p(full_result, reduced_result)
            except Exception:
                p_value = np.nan
            if pd.isna(p_value):
                continue
            if p_value > worst_p:
                worst_p = p_value
                worst_predictor = predictor

        if worst_predictor is None or worst_p < 0.05:
            break
        selected.remove(worst_predictor)

    if not selected:
        return pd.DataFrame(), pd.DataFrame(), []

    final_result, final_df = fit_multivariable_formula(data, outcome, selected, metadata)
    conf_int = final_result.conf_int()

    overall_rows = []
    for predictor in selected:
        reduced_predictors = [p for p in selected if p != predictor]
        if not reduced_predictors:
            overall_p = np.nan
        else:
            reduced_result, _ = fit_multivariable_formula(final_df, outcome, reduced_predictors, metadata)
            overall_p = likelihood_ratio_p(final_result, reduced_result)
        overall_rows.append(
            {
                "predictor": predictor,
                "变量": label_name(metadata[predictor]["source"]),
                "变量类型": metadata[predictor]["transform"],
                "整体P值": overall_p,
                "样本量": len(final_df),
            }
        )

    coef_rows = []
    for term in final_result.params.index:
        if term == "Intercept":
            continue
        matched_predictor = None
        for predictor in selected:
            if f'Q("{predictor}")' in term:
                matched_predictor = predictor
                break
        if matched_predictor is None:
            continue
        meta = metadata[matched_predictor]
        coef = float(final_result.params[term])
        se = float(final_result.bse[term])
        p_value = float(final_result.pvalues[term])
        lower = float(conf_int.loc[term, 0])
        upper = float(conf_int.loc[term, 1])
        variable_name, comparison = parse_term(term, matched_predictor, meta)
        overall_p = next(
            (row["整体P值"] for row in overall_rows if row["predictor"] == matched_predictor),
            np.nan,
        )
        coef_rows.append(
            {
                "变量": variable_name,
                "比较": comparison,
                "变量类型": meta["transform"],
                "回归系数β": coef,
                "标准误SE": se,
                "OR": safe_exp(coef),
                "OR 95%CI下限": safe_exp(lower),
                "OR 95%CI上限": safe_exp(upper),
                "P值": p_value,
                "整体P值": overall_p,
                "样本量": len(final_df),
            }
        )

    coef_df = pd.DataFrame(coef_rows).sort_values(["整体P值", "P值", "变量"], na_position="last")
    overall_df = pd.DataFrame(overall_rows).sort_values(["整体P值", "变量"], na_position="last")
    return coef_df.reset_index(drop=True), overall_df.reset_index(drop=True), selected


def display_table(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=["变量", "比较", "变量类型", "β", "SE", "OR", "95%CI", "P值", "整体P值", "样本量"])
    out = df.copy()
    out["β"] = out["回归系数β"].map(lambda x: fmt_num(x))
    out["SE"] = out["标准误SE"].map(lambda x: fmt_num(x))
    out["OR"] = out["OR"].map(lambda x: fmt_num(x))
    out["95%CI"] = out.apply(lambda r: f"{fmt_num(r['OR 95%CI下限'])} ~ {fmt_num(r['OR 95%CI上限'])}", axis=1)
    out["P值"] = out["P值"].map(lambda x: "<0.001" if x < 0.001 else f"{x:.3f}")
    out["整体P值"] = out["整体P值"].map(lambda x: "" if pd.isna(x) else ("<0.001" if x < 0.001 else f"{x:.3f}"))
    return out[["变量", "比较", "变量类型", "β", "SE", "OR", "95%CI", "P值", "整体P值", "样本量"]]


def display_publish_table(df: pd.DataFrame, significant_only: bool = False) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=["因素", "比较", "β", "SE", "OR(95%CI)", "P值"])
    out = df.copy()
    if significant_only:
        out = out[out["P值"] < 0.05].copy()
    if out.empty:
        return pd.DataFrame(columns=["因素", "比较", "β", "SE", "OR(95%CI)", "P值"])
    out["β"] = out["回归系数β"].map(fmt_num)
    out["SE"] = out["标准误SE"].map(fmt_num)
    out["OR(95%CI)"] = out.apply(
        lambda r: f"{fmt_num(r['OR'])} ({fmt_num(r['OR 95%CI下限'])}, {fmt_num(r['OR 95%CI上限'])})",
        axis=1,
    )
    out["P值"] = out["P值"].map(lambda x: "<0.001" if x < 0.001 else f"{x:.3f}")
    out = out.rename(columns={"变量": "因素"})
    return out[["因素", "比较", "β", "SE", "OR(95%CI)", "P值"]]


def prepare_forest_df(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=["标签", "OR", "OR 95%CI下限", "OR 95%CI上限", "P值"])
    out = df.copy()
    out["标签"] = out.apply(lambda r: f"{r['变量']}：{r['比较']}", axis=1)
    out = out[(out["OR"] > 0) & np.isfinite(out["OR"])].copy()
    out = out[(out["OR 95%CI下限"] > 0) & np.isfinite(out["OR 95%CI下限"])].copy()
    if out.empty:
        return pd.DataFrame(columns=["标签", "OR", "OR 95%CI下限", "OR 95%CI上限", "P值"])
    out["plot_upper"] = out["OR 95%CI上限"].replace(np.inf, np.nan)
    finite_upper = out["plot_upper"][np.isfinite(out["plot_upper"])]
    cap_upper = float(finite_upper.max()) if not finite_upper.empty else float(out["OR"].max() * 1.5)
    cap_upper = max(cap_upper, 1.5)
    out["plot_upper"] = out["plot_upper"].fillna(cap_upper * 1.2).clip(upper=cap_upper * 1.2)
    out = out.sort_values(["P值", "OR"], ascending=[True, False]).reset_index(drop=True)
    return out


def save_forest_plot(df: pd.DataFrame, title: str, out_file: Path):
    plot_df = prepare_forest_df(df)
    if plot_df.empty:
        return

    height = max(4.8, min(0.48 * len(plot_df) + 2.0, 16))
    fig, ax = plt.subplots(figsize=(10.8, height), facecolor="white")
    y_pos = np.arange(len(plot_df))[::-1]
    main_color = "#1D4ED8"
    sig_color = "#C2410C"
    neutral_color = "#475467"
    grid_color = "#D0D5DD"
    ci_color = "#94A3B8"

    lower_err = plot_df["OR"].values - plot_df["OR 95%CI下限"].values
    upper_err = plot_df["plot_upper"].values - plot_df["OR"].values
    colors = [sig_color if p < 0.05 else neutral_color for p in plot_df["P值"].values]

    ax.errorbar(
        plot_df["OR"].values,
        y_pos,
        xerr=[lower_err, upper_err],
        fmt="o",
        color="#111827",
        ecolor=ci_color,
        elinewidth=1.8,
        capsize=3.5,
        markersize=5.5,
        zorder=2,
    )
    for x, y, c in zip(plot_df["OR"].values, y_pos, colors):
        ax.scatter([x], [y], color=c, s=36, zorder=3, edgecolor="white", linewidth=0.6)

    ax.axvline(1.0, color="#DC2626", linestyle="--", linewidth=1.2, alpha=0.9)
    ax.set_xscale("log")
    ax.set_yticks(y_pos)
    ax.set_yticklabels(plot_df["标签"].tolist(), fontsize=9.5, color="#111827")
    ax.set_xlabel("OR（对数刻度）", fontsize=10.5, color="#111827", labelpad=10)
    ax.set_title(title, fontsize=13.5, weight="bold", color=main_color, pad=14)
    ax.grid(axis="x", linestyle=":", color=grid_color, alpha=0.8)
    ax.set_axisbelow(True)
    ax.set_facecolor("#F8FAFC")

    x_max = max(plot_df["plot_upper"].max(), plot_df["OR"].max(), 1.2)
    x_min = min(plot_df["OR 95%CI下限"].min(), 0.8)
    ax.set_xlim(max(x_min * 0.8, 0.01), x_max * 1.35)
    ax.tick_params(axis="x", labelsize=9, colors="#344054")
    ax.tick_params(axis="y", length=0)

    for _, row in plot_df.iterrows():
        y = y_pos[row.name]
        text = f"{fmt_num(row['OR'])} ({fmt_num(row['OR 95%CI下限'])}, {fmt_num(row['OR 95%CI上限'])})"
        color = sig_color if row["P值"] < 0.05 else "#334155"
        ax.text(x_max * 1.08, y, text, va="center", ha="left", fontsize=8.8, color=color)

    for spine in ["top", "right", "left"]:
        ax.spines[spine].set_visible(False)
    ax.spines["bottom"].set_color("#98A2B3")
    ax.text(
        0.995,
        1.02,
        "虚线表示 OR=1；橙色点表示 P<0.05",
        transform=ax.transAxes,
        ha="right",
        va="bottom",
        fontsize=8.5,
        color="#475467",
    )

    fig.tight_layout(pad=1.2)
    fig.savefig(out_file, dpi=300, bbox_inches="tight")
    plt.close(fig)


def format_three_line_sheet(ws, title: str):
    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    ws.insert_rows(1)
    ws["A1"] = title
    ws["A1"].font = Font(name="Microsoft YaHei", size=12, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ws.max_column)

    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cell in ws[2]:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True)
        cell.border = Border(top=medium, bottom=thin)

    if ws.max_row >= 3:
        for cell in ws[ws.max_row]:
            cell.border = Border(bottom=medium)

    for idx in range(1, ws.max_column + 1):
        letter = get_column_letter(idx)
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in ws[letter])
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 26)
    ws.freeze_panes = "A3"


def save_excel(file_path: Path, sheets: dict[str, pd.DataFrame]):
    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
        for sheet_name in writer.book.sheetnames:
            format_three_line_sheet(writer.book[sheet_name], sheet_name)


def build_summary(outcome_results: dict[str, dict], metadata: dict[str, dict]):
    lines = [
        "# 先心 ECMO Logistic 回归分析说明",
        "",
        "## 数据来源与分析假设",
        "- 数据文件：`先心ECMO临床资料.xlsx`",
        "- 主分析数据集：`Sheet2`",
        "- 理由：`Sheet2` 的结果变量和定量字段更规整，缺失率更低，更适合作为主分析表。",
        "- 结局变量：`撤机`、`出院`，分别整理为二分类结局：`是=1`、`否=0`。",
        "- 所有自变量缺失值在进入对应模型前直接剔除，不保留 `缺失` 组。",
        "- 定量变量先做正态性检验，若非正态则按三分位数分组后进入回归；缺失值同样剔除。",
        "",
        "## 变量处理方式",
    ]
    for safe_name, meta in metadata.items():
        lines.append(f"- `{label_name(meta['source'])}`：{meta['transform']}")

    for outcome_name, result in outcome_results.items():
        lines.extend(
            [
                "",
                f"## {outcome_name}",
                f"- 可用于分析的样本量：{result['n']}",
                f"- 单因素整体显著变量数：{result['uni_sig_count']}",
                f"- 多因素最终保留变量数：{len(result['selected_variables'])}",
            ]
        )
        if result["selected_variables"]:
            lines.append("- 多因素保留变量：")
            for variable in result["selected_variables"]:
                lines.append(f"  - {label_name(metadata[variable]['source'])}")
        else:
            lines.append("- 多因素模型未保留统计学显著变量。")

    lines.extend(
        [
            "",
            "## 输出文件",
            f"- `{PROCESSED_FILE.name}`：分析用清洗数据集",
            "- `withdraw_univariate.csv` / `withdraw_multivariate.csv`：撤机成功回归结果",
            "- `discharge_univariate.csv` / `discharge_multivariate.csv`：出院成功回归结果",
            "- `withdraw_group_comparison.csv` / `discharge_group_comparison.csv`：结局分组描述与组间检验表",
            f"- `{EXCEL_FILE.name}`：汇总 Excel",
            f"- `{PUBLISH_EXCEL_FILE.name}`：发表格式 Excel",
            "- `forest_plots/*.png`：OR 森林图",
        ]
    )
    SUMMARY_FILE.write_text("\n".join(lines), encoding="utf-8")


def build_result_narrative(outcome_results: dict[str, dict], metadata: dict[str, dict]):
    lines = [
        "# 回归分析结果描述",
        "",
        "## 方法概述",
        "- 数据来源：`先心ECMO临床资料.xlsx` 的 `Sheet2`。",
        "- 结局变量：`撤机成功`、`出院成功`，以“是”为1，“否”为0。",
        "- 所有自变量缺失值均不保留为单独分组，进入各模型前直接剔除。",
        "- 定量变量经正态性检验后，非正态变量按三分位数分组进入回归。",
        "- 先行单因素 Logistic 回归，再将单因素整体P<0.05的变量纳入多因素 Logistic 回归，采用逐步剔除法保留具有统计学意义的因素。",
    ]

    for outcome_name, result in outcome_results.items():
        lines.extend(["", f"## {outcome_name}"])
        uni = result["uni_overall"]
        sig_vars = uni.loc[uni["整体P值"] < 0.05, "变量"].tolist() if not uni.empty else []
        if sig_vars:
            lines.append(f"单因素分析显示，{('、'.join(sig_vars))}与{outcome_name}相关（整体P<0.05）。")
        else:
            lines.append(f"单因素分析未发现与{outcome_name}相关的显著变量。")
        multi = result["multi_coef"]
        selected = result["selected_variables"]
        if selected and not multi.empty:
            lines.append(f"多因素 Logistic 回归分析最终保留{len(selected)}个变量。")
            for _, row in multi.iterrows():
                p_text = "<0.001" if row["P值"] < 0.001 else f"{row['P值']:.3f}"
                lines.append(
                    f"- {row['变量']}{row['比较']}：β={row['回归系数β']:.3f}，SE={row['标准误SE']:.3f}，"
                    f"OR={row['OR']:.3f}，95%CI {row['OR 95%CI下限']:.3f}~{row['OR 95%CI上限']:.3f}，P={p_text}。"
                )
        else:
            lines.append("多因素 Logistic 回归分析未保留统计学显著变量。")

    lines.extend(
        [
            "",
            "## 结果解释提示",
            "- 以参考组为基准，OR>1 表示结局发生优势升高，OR<1 表示结局发生优势降低。",
            "- 对于按三分位数组的定量变量，比较基准为 Q1。",
        ]
    )
    RESULT_NARRATIVE_FILE.write_text("\n".join(lines), encoding="utf-8")


def main():
    analysis_df, metadata = build_analysis_dataset()
    descriptive_df, descriptive_meta = build_descriptive_dataset()
    analysis_df.to_csv(PROCESSED_FILE, index=False, encoding="utf-8-sig")

    excel_sheets = {}
    publish_sheets = {}
    outcome_results = {}

    for outcome_name in OUTCOMES:
        subset_n = int(analysis_df[outcome_name].notna().sum())
        uni_coef_df, uni_var_df = run_univariate(analysis_df, outcome_name, metadata)
        multi_coef_df, multi_var_df, selected = run_multivariate(analysis_df, outcome_name, metadata, uni_var_df)

        prefix = "withdraw" if outcome_name == "撤机成功" else "discharge"
        uni_coef_df.to_csv(RESULT_DIR / f"{prefix}_univariate.csv", index=False, encoding="utf-8-sig")
        multi_coef_df.to_csv(RESULT_DIR / f"{prefix}_multivariate.csv", index=False, encoding="utf-8-sig")
        uni_var_df.to_csv(RESULT_DIR / f"{prefix}_univariate_overall.csv", index=False, encoding="utf-8-sig")
        multi_var_df.to_csv(RESULT_DIR / f"{prefix}_multivariate_overall.csv", index=False, encoding="utf-8-sig")
        group_table_df = build_group_comparison_table(descriptive_df, outcome_name, descriptive_meta)
        group_table_df.to_csv(RESULT_DIR / f"{prefix}_group_comparison.csv", index=False, encoding="utf-8-sig")

        excel_sheets[f"{outcome_name}-单因素"] = display_table(uni_coef_df)
        excel_sheets[f"{outcome_name}-多因素"] = display_table(multi_coef_df)
        excel_sheets[f"{outcome_name}-组间描述"] = group_table_df
        publish_sheets[f"{outcome_name}-单因素表"] = display_publish_table(uni_coef_df)
        publish_sheets[f"{outcome_name}-多因素表"] = display_publish_table(multi_coef_df)
        publish_sheets[f"{outcome_name}-多因素显著项"] = display_publish_table(multi_coef_df, significant_only=True)
        publish_sheets[f"{outcome_name}-组间描述表"] = group_table_df

        forest_prefix = "withdraw" if outcome_name == "撤机成功" else "discharge"
        save_forest_plot(uni_coef_df, f"{outcome_name}：单因素 Logistic 回归 OR 森林图", FOREST_DIR / f"{forest_prefix}_univariate_forest.png")
        save_forest_plot(multi_coef_df, f"{outcome_name}：多因素 Logistic 回归 OR 森林图", FOREST_DIR / f"{forest_prefix}_multivariate_forest.png")

        outcome_results[outcome_name] = {
            "n": subset_n,
            "uni_sig_count": int((uni_var_df["整体P值"] < 0.05).sum()) if not uni_var_df.empty else 0,
            "selected_variables": selected,
            "uni_overall": uni_var_df,
            "multi_coef": multi_coef_df,
        }

    save_excel(EXCEL_FILE, excel_sheets)
    save_excel(PUBLISH_EXCEL_FILE, publish_sheets)
    build_summary(outcome_results, metadata)
    build_result_narrative(outcome_results, metadata)


if __name__ == "__main__":
    main()
